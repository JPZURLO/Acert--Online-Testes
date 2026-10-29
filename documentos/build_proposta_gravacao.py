from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "documentos" / "Proposta_Gravacao_e_Auditoria_de_Testes.docx"
LOGO = ROOT / "front-end" / "assets" / "images" / "Logo.png"

# narrative_proposal preset + named brand override: petroleum accent.
PETROLEUM = "0F6F73"
DARK = "26373A"
MUTED = "64777A"
LIGHT = "F2F7F7"
LIGHTER = "F8FBFB"
BLUE_GRAY = "E8EEF5"
GREEN = "16815E"
AMBER = "A66A00"
RED = "A33A42"
WHITE = "FFFFFF"
BORDER = "C9D7D8"


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


def set_font(run, size=11, bold=False, italic=False, color=DARK, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = rgb(color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
    first_row_pr = table.rows[0]._tr.get_or_add_trPr()
    header_marker = first_row_pr.find(qn("w:tblHeader"))
    if header_marker is None:
        header_marker = OxmlElement("w:tblHeader")
        first_row_pr.append(header_marker)
    header_marker.set(qn("w:val"), "true")


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    set_font(run, 9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), PETROLEUM)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_pr.extend([color, underline])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([run_pr, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def configure_numbering(document):
    numbering = document.part.numbering_part.element
    abstract_id = 42
    num_id = 42
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet")
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "•")
    level_jc = OxmlElement("w:lvlJc")
    level_jc.set(qn("w:val"), "left")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "280")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "290")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.extend([tabs, ind, spacing])
    level.extend([start, num_fmt, level_text, level_jc, p_pr])
    abstract.append(level)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_bullet(document, text, num_id):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.208
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_node])
    p_pr.append(num_pr)
    set_font(paragraph.add_run(text), 11)
    return paragraph


def add_heading(document, text, level=1):
    paragraph = document.add_paragraph(style=f"Heading {level}")
    paragraph.add_run(text)
    keep_with_next(paragraph)
    return paragraph


def add_body(document, text, bold_lead=None):
    paragraph = document.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        set_font(paragraph.add_run(bold_lead), 11, bold=True)
        set_font(paragraph.add_run(text[len(bold_lead):]), 11)
    else:
        set_font(paragraph.add_run(text), 11)
    return paragraph


def add_callout(document, label, text, fill=LIGHT, accent=PETROLEUM):
    table = document.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_table_borders(table, color=accent, size="8")
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    set_font(paragraph.add_run(f"{label} "), 11, bold=True, color=accent)
    set_font(paragraph.add_run(text), 11, color=DARK)
    after = document.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    return table


def format_table_text(table, header=True):
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.08
                for run in paragraph.runs:
                    set_font(run, 9.5, bold=header and row_index == 0, color=DARK)
            if header and row_index == 0:
                shade_cell(cell, BLUE_GRAY)


def configure_document():
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(DARK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    heading_tokens = {
        1: (16, 18, 10),
        2: (13, 12, 6),
        3: (12, 8, 4),
    }
    for level, (size, before, after) in heading_tokens.items():
        style = document.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(PETROLEUM)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    header_paragraph = header.paragraphs[0]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_font(header_paragraph.add_run("ONLINE TESTE  |  PROPOSTA DE GRAVAÇÃO E AUDITORIA"), 8.5, bold=True, color=MUTED)
    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    add_page_number(footer_paragraph)
    return document


def build():
    document = configure_document()
    bullet_id = configure_numbering(document)

    # Cover: proposal_centerpiece pattern.
    if LOGO.exists():
        logo_p = document.add_paragraph()
        logo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        logo_p.paragraph_format.space_before = Pt(18)
        logo_p.paragraph_format.space_after = Pt(30)
        logo_shape = logo_p.add_run().add_picture(str(LOGO), width=Inches(1.7))
        logo_shape._inline.docPr.set("descr", "Logotipo Online Teste")
        logo_shape._inline.docPr.set("title", "Online Teste")
    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(10)
    set_font(kicker.add_run("PROPOSTA TÉCNICA E COMERCIAL"), 10, bold=True, color=PETROLEUM)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    set_font(title.add_run("Gravação e auditoria\nde testes online"), 28, bold=True, color=DARK)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(30)
    set_font(subtitle.add_run("Comparativo entre implantação gratuita e infraestrutura profissional paga"), 14, color=MUTED)
    add_callout(
        document,
        "OBJETIVO",
        "Apresentar caminhos viáveis para registrar câmera, microfone e tela do participante, permitindo auditoria das aplicações sem comprometer a evolução futura da plataforma.",
        fill=LIGHTER,
    )
    metadata = document.add_table(rows=3, cols=2)
    metadata_data = [
        ("Projeto", "Online Teste"),
        ("Documento", "Análise de opções de gravação e fiscalização"),
        ("Data-base", "Julho de 2026"),
    ]
    for row, values in zip(metadata.rows, metadata_data):
        for cell, value in zip(row.cells, values):
            cell.text = value
    set_table_geometry(metadata, [2200, 7160])
    set_table_borders(metadata, color=BORDER, size="4")
    for row in metadata.rows:
        shade_cell(row.cells[0], LIGHT)
        set_font(row.cells[0].paragraphs[0].runs[0], 9.5, bold=True, color=PETROLEUM)
        set_font(row.cells[1].paragraphs[0].runs[0], 9.5, color=DARK)
    document.add_page_break()

    add_heading(document, "1. Resumo executivo", 1)
    add_body(document, "A plataforma pode iniciar com uma modalidade sem nova mensalidade, utilizando gravação pelo navegador e armazenamento gratuito limitado. Essa alternativa é adequada para validação comercial e baixo volume, mas possui menor robustez de auditoria.")
    add_body(document, "Para uma operação profissional, a recomendação é transmitir a mídia para um serviço de gravação no servidor e armazenar os arquivos em ambiente imutável. Essa estrutura reduz o risco de perda ou manipulação, permite fiscalização ao vivo e suporta crescimento.")
    add_callout(document, "RECOMENDAÇÃO", "Começar pela opção gratuita, porém desenvolver desde o início uma camada substituível de gravação e armazenamento. Quando houver clientes pagantes, a plataforma migra para a estrutura profissional sem reconstruir as telas e regras de negócio.")

    add_heading(document, "2. Comparativo das opções", 1)
    table = document.add_table(rows=1, cols=3)
    headers = ["Critério", "Opção gratuita", "Opção paga profissional"]
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    comparison = [
        ("Custo recorrente", "R$ 0 dentro das franquias", "A partir de cerca de R$ 300/mês"),
        ("Captura", "Navegador do participante", "Transmissão e gravação no servidor"),
        ("Armazenamento", "10 a 20 GB gratuitos", "Escalável e com retenção configurável"),
        ("Risco de perda", "Pode perder o último bloco em falha brusca", "Muito menor; gravação ocorre fora do computador"),
        ("Integridade", "Hashes e registros demonstram alterações", "Pode usar armazenamento imutável"),
        ("Acompanhamento ao vivo", "Não incluído inicialmente", "Disponível no mesmo fluxo"),
        ("Escala", "Baixo volume", "Múltiplas aplicações simultâneas"),
        ("Indicação", "Piloto e validação", "Operação comercial e auditoria formal"),
    ]
    for values in comparison:
        row = table.add_row()
        for cell, value in zip(row.cells, values):
            cell.text = value
    set_table_geometry(table, [1900, 3450, 4010])
    set_table_borders(table)
    format_table_text(table)

    add_heading(document, "3. Opção gratuita: piloto operacional", 1)
    add_body(document, "Nesta modalidade, o navegador grava a câmera com o áudio e, separadamente, a tela inteira. Os dados são enviados continuamente ao servidor e arquivados em uma franquia gratuita de nuvem.")
    add_heading(document, "Como funciona", 2)
    for item in [
        "O participante envia documento e selfie e autoriza câmera, microfone e compartilhamento da tela inteira.",
        "A gravação é dividida em fragmentos contínuos de aproximadamente cinco segundos.",
        "Cada fragmento recebe número sequencial, horário e hash SHA-256.",
        "O navegador mantém cópia temporária até receber a confirmação do servidor.",
        "Interromper câmera, áudio, tela ou sair do ambiente seguro encerra a prova.",
        "Ao concluir, o servidor consolida os fragmentos e envia os vídeos ao armazenamento escolhido.",
    ]:
        add_bullet(document, item, bullet_id)

    add_heading(document, "Pontos positivos", 2)
    for item in [
        "Permite iniciar sem contratar uma nova mensalidade.",
        "Valida o comportamento dos participantes e o volume real de gravações.",
        "A estrutura de hashes e eventos melhora a rastreabilidade.",
        "Pode migrar futuramente para outro provedor sem mudar a experiência do usuário.",
    ]:
        add_bullet(document, item, bullet_id)

    add_heading(document, "Limitações", 2)
    for item in [
        "As franquias gratuitas comportam apenas poucas dezenas de horas de vídeo.",
        "Uma falha brusca pode causar a perda do fragmento ainda não entregue.",
        "Os arquivos não têm a mesma proteção contra exclusão de um armazenamento imutável.",
        "A fiscalização ao vivo exige infraestrutura adicional.",
    ]:
        add_bullet(document, item, bullet_id)
    document.add_page_break()

    add_heading(document, "4. Segurança dos fragmentos de gravação", 1)
    add_body(document, "A divisão em fragmentos não cria lacunas quando existe uma única sessão contínua do MediaRecorder. O intervalo solicitado ao navegador, entretanto, não é um relógio exato e pode atrasar em determinadas condições.")
    add_callout(document, "RISCO RESIDUAL", "Em desligamento, travamento ou encerramento forçado, pode ser perdido o fragmento que ainda estava apenas na memória. Com fragmentos de cinco segundos, a exposição normalmente se limita aos segundos finais, embora atrasos do navegador possam ampliar esse período.", fill="FFF7E8", accent=AMBER)
    add_heading(document, "Controles necessários", 2)
    controls = [
        ("Sequência", "O servidor rejeita duplicidades e identifica números ausentes."),
        ("Confirmação", "O navegador só remove o fragmento local após confirmação do servidor."),
        ("Fila offline", "Fragmentos pendentes ficam no IndexedDB e são reenviados quando a conexão retorna."),
        ("Integridade", "Cada fragmento e o arquivo final recebem hash SHA-256."),
        ("Finalização", "A prova só é concluída após o envio e a confirmação do último fragmento."),
        ("Consolidação", "O servidor usa FFmpeg para remontar os arquivos na ordem correta."),
    ]
    control_table = document.add_table(rows=1, cols=2)
    control_table.rows[0].cells[0].text = "Controle"
    control_table.rows[0].cells[1].text = "Aplicação"
    for label, detail in controls:
        row = control_table.add_row()
        row.cells[0].text = label
        row.cells[1].text = detail
    set_table_geometry(control_table, [2300, 7060])
    set_table_borders(control_table)
    format_table_text(control_table)

    add_heading(document, "5. Armazenamento gratuito e econômico", 1)
    storage = document.add_table(rows=1, cols=4)
    for cell, value in zip(storage.rows[0].cells, ["Serviço", "Franquia", "Uso recomendado", "Observação"]):
        cell.text = value
    rows = [
        ("MEGA", "20 GB", "Piloto muito pequeno", "Maior espaço gratuito, mas integração de auditoria menos conveniente."),
        ("Backblaze B2", "10 GB", "Melhor opção econômica", "Após a franquia: cerca de US$ 0,00695 por GB/mês."),
        ("Cloudflare R2", "10 GB", "Vídeos acessados com frequência", "Sem cobrança de saída; cerca de US$ 0,015 por GB/mês."),
        ("Google Drive", "Conforme a conta", "Aproveitar espaço já contratado", "API padrão sem custo adicional dentro das quotas atuais."),
    ]
    for values in rows:
        row = storage.add_row()
        for cell, value in zip(row.cells, values):
            cell.text = value
    set_table_geometry(storage, [1450, 1350, 2440, 4120])
    set_table_borders(storage)
    format_table_text(storage)
    add_body(document, "Com gravações estimadas entre 1 e 1,5 GB por hora, 20 GB armazenam aproximadamente 13 a 20 horas. Por isso, a opção gratuita exige retenção curta e exclusão automática.")
    document.add_page_break()

    add_heading(document, "6. Opção paga: auditoria profissional", 1)
    add_body(document, "Na modalidade profissional, câmera, microfone e tela são enviados por WebRTC para um serviço de mídia. A gravação ocorre no servidor, independentemente do arquivo local do participante.")
    add_heading(document, "Arquitetura sugerida", 2)
    architecture = document.add_table(rows=1, cols=5)
    for index, value in enumerate(["Participante", "WebRTC", "LiveKit", "Armazenamento", "Auditoria"]):
        architecture.rows[0].cells[index].text = value
    row = architecture.add_row()
    for index, value in enumerate(["Câmera, áudio e tela", "Transmissão segura", "Gravação no servidor", "S3 Object Lock", "Player e eventos"]):
        row.cells[index].text = value
    set_table_geometry(architecture, [1740, 1500, 1800, 2200, 2120])
    set_table_borders(architecture, color=PETROLEUM, size="8")
    format_table_text(architecture)
    for cell in architecture.rows[0].cells:
        shade_cell(cell, LIGHT)

    add_heading(document, "Benefícios", 2)
    for item in [
        "Menor risco de perda e manipulação das gravações.",
        "Acompanhamento ao vivo opcional sem controle remoto do computador.",
        "Armazenamento com retenção e bloqueio contra alteração ou exclusão.",
        "Escalabilidade para múltiplos participantes e empresas.",
        "Registro de visualizações, downloads, alertas e decisões do fiscal.",
    ]:
        add_bullet(document, item, bullet_id)

    add_heading(document, "Estimativa de custos", 2)
    costs = document.add_table(rows=1, cols=3)
    for cell, value in zip(costs.rows[0].cells, ["Componente", "Referência", "Estimativa"]):
        cell.text = value
    cost_rows = [
        ("LiveKit Build", "Desenvolvimento e testes", "US$ 0/mês, com franquias reduzidas"),
        ("LiveKit Ship", "Produção", "US$ 50/mês"),
        ("Vídeo transcodificado", "Após a franquia", "US$ 0,02 por minuto"),
        ("Faixa original", "Após a franquia", "US$ 0,001 por faixa/minuto"),
        ("Produção de baixo volume", "LiveKit + armazenamento", "Aproximadamente R$ 300 a R$ 600/mês"),
    ]
    for values in cost_rows:
        row = costs.add_row()
        for cell, value in zip(row.cells, values):
            cell.text = value
    set_table_geometry(costs, [2700, 2600, 4060])
    set_table_borders(costs)
    format_table_text(costs)
    add_body(document, "Os valores são estimativas de planejamento, considerando câmbio próximo de R$ 5,15 por dólar. Impostos, câmbio, volume, resolução, downloads e retenção podem alterar o total.")

    add_heading(document, "7. Recomendação de implantação", 1)
    add_callout(document, "DECISÃO SUGERIDA", "Aprovar a implantação gratuita como piloto, com gravação em fragmentos, hashes, fila offline e armazenamento substituível. A opção profissional deve permanecer prevista para ativação quando houver clientes pagantes ou exigência contratual de auditoria formal.", fill=LIGHT, accent=GREEN)
    roadmap = document.add_table(rows=1, cols=3)
    for cell, value in zip(roadmap.rows[0].cells, ["Etapa", "Entrega", "Condição"]):
        cell.text = value
    roadmap_rows = [
        ("1. Piloto", "Gravação local, fragmentos e nuvem gratuita", "Sem nova mensalidade"),
        ("2. Validação", "Player de auditoria, retenção e relatórios", "Primeiros clientes"),
        ("3. Profissional", "LiveKit, gravação no servidor e S3", "Receita recorrente"),
        ("4. Ao vivo", "Painel de fiscais e alertas em tempo real", "Demanda comercial comprovada"),
    ]
    for values in roadmap_rows:
        row = roadmap.add_row()
        for cell, value in zip(row.cells, values):
            cell.text = value
    set_table_geometry(roadmap, [1700, 4800, 2860])
    set_table_borders(roadmap)
    format_table_text(roadmap)

    add_heading(document, "8. Texto resumido para apresentação", 1)
    add_callout(
        document,
        "MENSAGEM",
        "Podemos iniciar sem custo adicional, utilizando gravação em fragmentos e armazenamento gratuito. Essa modalidade atende à validação inicial, mas possui limites de espaço e menor força de auditoria. Quando o volume ou a exigência de segurança aumentar, a mesma plataforma poderá migrar para gravação no servidor, armazenamento imutável e acompanhamento ao vivo, com custo estimado a partir de R$ 300 mensais.",
        fill=LIGHTER,
    )

    add_heading(document, "Fontes consultadas", 2)
    sources = [
        ("LiveKit - preços", "https://livekit.com/pricing"),
        ("Amazon S3 - Object Lock", "https://docs.aws.amazon.com/AmazonS3/latest/userguide/object-lock.html"),
        ("Backblaze B2 - preços", "https://www.backblaze.com/cloud-storage/transaction-pricing"),
        ("Cloudflare R2 - preços", "https://developers.cloudflare.com/r2/pricing/"),
        ("Google Drive API - limites", "https://developers.google.com/workspace/drive/api/guides/limits"),
        ("MDN - MediaRecorder", "https://developer.mozilla.org/en-US/docs/Web/API/MediaRecorder/dataavailable_event"),
    ]
    for label, url in sources:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(3)
        set_font(paragraph.add_run(f"{label}: "), 9, bold=True, color=DARK)
        add_hyperlink(paragraph, "Acessar fonte", url)

    document.core_properties.title = "Proposta de gravação e auditoria de testes online"
    document.core_properties.subject = "Comparativo entre implantação gratuita e opção profissional paga"
    document.core_properties.author = "Online Teste"
    document.core_properties.keywords = "gravação, auditoria, testes online, LiveKit, armazenamento"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
