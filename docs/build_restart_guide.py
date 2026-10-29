from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "Guia_para_reativar_site_de_homologacao.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "24313D"
MUTED = "66717D"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F4F6F9"
GREEN = "16735B"
CAUTION = "FFF4CE"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_run_font(run, name="Calibri", size=11, color=INK, bold=False, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[idx]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")

            tc_mar = tc_pr.find(qn("w:tcMar"))
            if tc_mar is None:
                tc_mar = OxmlElement("w:tcMar")
                tc_pr.append(tc_mar)
            for edge, value in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
                node = tc_mar.find(qn(f"w:{edge}"))
                if node is None:
                    node = OxmlElement(f"w:{edge}")
                    tc_mar.append(node)
                node.set(qn("w:w"), str(value))
                node.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    set_run_font(run, size=9, color=MUTED)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_numbering_definition(doc):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1.")
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(tabs)
    p_pr.append(ind)
    p_pr.append(spacing)
    lvl.extend([start, num_fmt, lvl_text, suff, lvl_jc, p_pr])
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abs_id = OxmlElement("w:abstractNumId")
    abs_id.set(qn("w:val"), str(abstract_id))
    num.append(abs_id)
    numbering.append(num)
    return num_id


def add_numbered_step(doc, num_id, title, detail):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_node])
    p_pr.append(num_pr)
    run = p.add_run(title)
    set_run_font(run, size=11, color=DARK_BLUE, bold=True)
    run = p.add_run(f" — {detail}")
    set_run_font(run, size=11, color=INK)


def add_code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360], 120)
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_GRAY)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    for idx, line in enumerate(lines):
        if idx:
            p.add_run().add_break()
        run = p.add_run(line)
        set_run_font(run, name="Consolas", size=9.2, color="17212B")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    spacer.paragraph_format.line_spacing = 0.5


def add_note(doc, label, text, fill=LIGHT_BLUE, color=DARK_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360], 120)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"{label}: ")
    set_run_font(r, size=10.5, color=color, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_label_paragraph(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(f"{label}: ")
    set_run_font(r, bold=True, color=DARK_BLUE)
    r = p.add_run(text)
    set_run_font(r)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.82)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    configure_styles(doc)
    num_id = add_numbering_definition(doc)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("ACERT  |  Ambiente de homologação")
    set_run_font(r, size=9, color=MUTED, bold=True)
    add_page_number(section.footer.paragraphs[0])

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(2)
    r = kicker.add_run("GUIA RÁPIDO")
    set_run_font(r, size=10, color=GREEN, bold=True)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)
    r = title.add_run("Como reativar o site para testes externos")
    set_run_font(r, size=25, color=DARK_BLUE, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    r = subtitle.add_run("Servidor local + banco de homologação + link temporário do Cloudflare Tunnel")
    set_run_font(r, size=12, color=MUTED)

    add_note(
        doc,
        "Resumo",
        "Sempre serão usadas duas janelas do PowerShell: uma mantém o site funcionando e a outra mantém o link público ativo.",
        fill=LIGHT_BLUE,
    )

    doc.add_heading("Antes de começar", level=1)
    add_label_paragraph(doc, "Pasta do projeto", r"H:\Meu Drive\online\Acert--Online-Testes")
    add_label_paragraph(doc, "Banco para os testes", "acert_homologacao")
    add_label_paragraph(doc, "Porta do site", "5500")
    add_label_paragraph(doc, "Programa do túnel", r"C:\Users\usr030\Downloads\cloudflared-windows-amd64.exe")

    doc.add_heading("Ativação passo a passo", level=1)
    add_numbered_step(doc, num_id, "Confirme se o MySQL está funcionando", "abra o PowerShell como administrador e execute:")
    add_code_block(doc, ["Get-Service MySQL80", "Start-Service MySQL80"])
    add_note(doc, "Observação", "Se o serviço já aparecer como Running, não há problema: siga para o próximo passo.")

    add_numbered_step(doc, num_id, "Abra o primeiro PowerShell", "ele ficará responsável pelo servidor do site.")
    add_code_block(
        doc,
        [
            'cd "H:\\Meu Drive\\online\\Acert--Online-Testes"',
            ".\\.venv\\Scripts\\Activate.ps1",
            '$env:DB_NAME="acert_homologacao"',
            '$env:COOKIE_SECURE="true"',
            '$env:PRIVATE_UPLOAD_DIR=(Join-Path $PWD "private_uploads_homologacao")',
            '$env:PRIVATE_RECORDING_DIR=(Join-Path $PWD "private_uploads_homologacao\\recordings")',
            "python server.py",
        ],
    )
    add_note(doc, "Importante", "Não feche essa janela. Se ela for fechada, o site para de funcionar.", fill=CAUTION)

    add_numbered_step(doc, num_id, "Abra um segundo PowerShell", "ele criará o endereço acessível por outras redes.")
    add_code_block(
        doc,
        [
            '& "$env:USERPROFILE\\Downloads\\cloudflared-windows-amd64.exe" tunnel --url http://localhost:5500'
        ],
    )

    add_numbered_step(doc, num_id, "Copie o link exibido", "procure a mensagem “Your quick Tunnel has been created”.")
    add_code_block(doc, ["https://nome-aleatorio.trycloudflare.com"])
    add_note(
        doc,
        "Atenção",
        "O endereço muda sempre que o Cloudflare Tunnel é encerrado e iniciado novamente. Use o link novo a cada sessão.",
        fill=CAUTION,
    )

    add_numbered_step(doc, num_id, "Teste antes de enviar", "abra o link no celular usando a rede móvel ou em outro computador.")
    add_numbered_step(doc, num_id, "Compartilhe somente com os validadores", "o endereço deixa temporariamente o ambiente de homologação acessível pela internet.")

    doc.add_heading("Durante os testes", level=1)
    add_label_paragraph(doc, "Janela 1", "deve continuar mostrando o servidor Python em execução.")
    add_label_paragraph(doc, "Janela 2", "deve continuar mostrando o Cloudflare Tunnel conectado.")
    add_label_paragraph(doc, "Computador", "precisa permanecer ligado e conectado à internet.")
    add_label_paragraph(doc, "MySQL", "o serviço MySQL80 precisa permanecer em execução.")

    doc.add_heading("Como encerrar", level=1)
    p = doc.add_paragraph()
    r = p.add_run("Em cada janela do PowerShell, pressione ")
    set_run_font(r)
    r = p.add_run("Ctrl + C")
    set_run_font(r, name="Consolas", color=DARK_BLUE, bold=True)
    r = p.add_run(". Depois disso, o link público deixa de funcionar.")
    set_run_font(r)

    doc.add_heading("Solução rápida de problemas", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [2700, 6660], 120)
    table.rows[0].cells[0].text = "Problema"
    table.rows[0].cells[1].text = "O que verificar"
    set_repeat_table_header(table.rows[0])
    for cell in table.rows[0].cells:
        set_cell_shading(cell, LIGHT_BLUE)
        for run in cell.paragraphs[0].runs:
            set_run_font(run, size=10, color=DARK_BLUE, bold=True)

    rows = [
        ("Erro 502 / Bad Gateway", "Confirme se o primeiro PowerShell ainda está aberto e se python server.py está rodando na porta 5500."),
        ("Acesso negado ao banco", "Confirme DB_NAME=acert_homologacao e as credenciais DB_USER/DB_PASSWORD do arquivo .env."),
        ("cloudflared não encontrado", "Confira se cloudflared-windows-amd64.exe continua na pasta Downloads."),
        ("Link antigo não abre", "Gere um novo túnel. O endereço trycloudflare.com muda depois que a janela é fechada."),
        ("MySQL está parado", "Abra o PowerShell como administrador e execute Start-Service MySQL80."),
    ]
    for problem, solution in rows:
        cells = table.add_row().cells
        cells[0].text = problem
        cells[1].text = solution
        for idx, cell in enumerate(cells):
            for para in cell.paragraphs:
                para.paragraph_format.space_before = Pt(0)
                para.paragraph_format.space_after = Pt(2)
                para.paragraph_format.line_spacing = 1.08
                for run in para.runs:
                    set_run_font(run, size=9.5, color=INK, bold=(idx == 0))
    set_table_geometry(table, [2700, 6660], 120)

    doc.add_heading("Checklist antes de mandar o link", level=1)
    checklist = [
        "MySQL80 aparece como Running.",
        "O servidor Python está aberto usando acert_homologacao.",
        "O Cloudflare Tunnel exibe Registered tunnel connection.",
        "O link abre fora da sua rede local.",
        "Somente dados de teste estão sendo usados.",
    ]
    for item in checklist:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run("☐  ")
        set_run_font(r, size=12, color=GREEN)
        r = p.add_run(item)
        set_run_font(r, size=10.5)

    add_note(
        doc,
        "Segurança",
        "Use somente o banco de homologação e dados fictícios. Para operação definitiva, será necessário publicar o sistema em uma hospedagem estável com domínio próprio, HTTPS e controles permanentes de acesso.",
        fill=CAUTION,
    )

    doc.core_properties.title = "Guia para reativar o site de homologação"
    doc.core_properties.subject = "Servidor local e Cloudflare Tunnel"
    doc.core_properties.author = "ACERT"
    doc.core_properties.keywords = "homologação, Cloudflare Tunnel, PowerShell, MySQL"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
